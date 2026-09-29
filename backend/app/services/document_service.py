"""
document_service.py

Generates downloadable PDF and DOCX documents from a structured MOM
(MOMResult). Both formats follow the same content structure:
title/date, summary, participants, discussion points, decisions,
action items (as a table), pending issues.

Empty sections show an explicit "none found" message rather than a
blank gap - this keeps the hallucination-control principle visible
in the exported document too, not just the app UI (Section 9: an
empty list means the AI found nothing, not that something was left
out by mistake).
"""

import io

from docx import Document
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import (
    ListFlowable,
    ListItem,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

from app.schemas.mom import MOMResult

ACCENT_COLOR = colors.HexColor("#2F6F5E")


def generate_pdf(mom: MOMResult) -> bytes:
    """Render a MOM as a PDF and return the raw file bytes."""
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=letter,
        topMargin=0.75 * inch,
        bottomMargin=0.75 * inch,
        leftMargin=0.75 * inch,
        rightMargin=0.75 * inch,
    )
    styles = getSampleStyleSheet()
    story = []

    title_text = mom.meeting_title or "Minutes of Meeting"
    story.append(Paragraph(title_text, styles["Title"]))
    if mom.meeting_date:
        story.append(Paragraph(f"Date: {mom.meeting_date}", styles["Normal"]))
    story.append(Spacer(1, 16))

    story.append(Paragraph("Summary", styles["Heading2"]))
    story.append(Paragraph(mom.summary or "No summary available.", styles["Normal"]))
    story.append(Spacer(1, 14))

    def add_bullet_section(heading: str, items: list[str], empty_text: str):
        story.append(Paragraph(heading, styles["Heading2"]))
        if items:
            bullets = ListFlowable(
                [ListItem(Paragraph(item, styles["Normal"])) for item in items],
                bulletType="bullet",
            )
            story.append(bullets)
        else:
            story.append(Paragraph(f"<i>{empty_text}</i>", styles["Normal"]))
        story.append(Spacer(1, 14))

    add_bullet_section("Participants", mom.participants, "Not specified.")
    add_bullet_section("Discussion Points", mom.discussion_points, "None recorded.")
    add_bullet_section("Decisions", mom.decisions, "No clear decisions were made.")

    story.append(Paragraph("Action Items", styles["Heading2"]))
    if mom.action_items:
        table_data = [["Task", "Assignee", "Deadline", "Status", "Priority"]]
        for item in mom.action_items:
            table_data.append(
                [
                    Paragraph(item.task, styles["Normal"]),
                    item.assignee or "Not specified",
                    item.deadline or "Not specified",
                    item.status or "Not specified",
                    item.priority or "Not specified",
                ]
            )
        table = Table(
            table_data,
            colWidths=[2.8 * inch, 1.2 * inch, 1.1 * inch, 1.0 * inch, 0.9 * inch],
        )
        table.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), ACCENT_COLOR),
                    ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                    ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                    ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F7F8FA")]),
                    ("TOPPADDING", (0, 0), (-1, -1), 6),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
                ]
            )
        )
        story.append(table)
    else:
        story.append(Paragraph("<i>No action items were identified.</i>", styles["Normal"]))
    story.append(Spacer(1, 14))

    add_bullet_section("Pending Issues", mom.pending_issues, "None.")

    doc.build(story)
    buffer.seek(0)
    return buffer.getvalue()


def generate_docx(mom: MOMResult) -> bytes:
    """Render a MOM as a DOCX and return the raw file bytes."""
    document = Document()

    title_text = mom.meeting_title or "Minutes of Meeting"
    document.add_heading(title_text, level=0)
    if mom.meeting_date:
        document.add_paragraph(f"Date: {mom.meeting_date}")

    document.add_heading("Summary", level=1)
    document.add_paragraph(mom.summary or "No summary available.")

    def add_bullet_section(heading: str, items: list[str], empty_text: str):
        document.add_heading(heading, level=1)
        if items:
            for item in items:
                document.add_paragraph(item, style="List Bullet")
        else:
            paragraph = document.add_paragraph()
            run = paragraph.add_run(empty_text)
            run.italic = True

    add_bullet_section("Participants", mom.participants, "Not specified.")
    add_bullet_section("Discussion Points", mom.discussion_points, "None recorded.")
    add_bullet_section("Decisions", mom.decisions, "No clear decisions were made.")

    document.add_heading("Action Items", level=1)
    if mom.action_items:
        table = document.add_table(rows=1, cols=5)
        table.style = "Light Grid Accent 1"
        header_cells = table.rows[0].cells
        header_cells[0].text = "Task"
        header_cells[1].text = "Assignee"
        header_cells[2].text = "Deadline"
        header_cells[3].text = "Status"
        header_cells[4].text = "Priority"
        for item in mom.action_items:
            row_cells = table.add_row().cells
            row_cells[0].text = item.task
            row_cells[1].text = item.assignee or "Not specified"
            row_cells[2].text = item.deadline or "Not specified"
            row_cells[3].text = item.status or "Not specified"
            row_cells[4].text = item.priority or "Not specified"
    else:
        paragraph = document.add_paragraph()
        run = paragraph.add_run("No action items were identified.")
        run.italic = True

    add_bullet_section("Pending Issues", mom.pending_issues, "None.")

    buffer = io.BytesIO()
    document.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()
